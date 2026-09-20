import io

from pptx import Presentation

from agents.prospection.offer import (
    default_highlights_for_pack,
    generate_offer_pdf,
    generate_offer_pptx,
)
from platform_core.models import Pack


def test_generate_offer_pdf_returns_valid_pdf_bytes() -> None:
    pdf_bytes = generate_offer_pdf("Chez Awa", Pack.business, ["Site vitrine", "Réseaux sociaux"])

    assert pdf_bytes.startswith(b"%PDF")
    assert len(pdf_bytes) > 100


def test_default_highlights_differ_by_pack() -> None:
    assert default_highlights_for_pack(Pack.starter) != default_highlights_for_pack(Pack.premium)


def test_generate_offer_pptx_returns_valid_presentation() -> None:
    pptx_bytes = generate_offer_pptx("Chez Awa", Pack.business, ["Site vitrine", "Réseaux sociaux"])

    # Format Office Open XML : un fichier .pptx est une archive zip, signature "PK".
    assert pptx_bytes.startswith(b"PK")

    presentation = Presentation(io.BytesIO(pptx_bytes))
    assert len(presentation.slides) == 1

    all_text = "\n".join(
        paragraph.text
        for shape in presentation.slides[0].shapes
        if shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
    )
    assert "Chez Awa" in all_text
    assert "Business" in all_text
    assert "Site vitrine" in all_text
    assert "Réseaux sociaux" in all_text

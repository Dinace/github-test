"""Génération d'un support visuel d'offre (PDF ou PowerPoint) pour un prospect, par pack."""

import io

from pptx import Presentation
from pptx.util import Cm, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas

from platform_core.models import Pack

_PACK_LABELS = {
    Pack.starter: "Starter",
    Pack.business: "Business",
    Pack.premium: "Premium",
}

# Résumé court par pack (CLAUDE.md §1, table des packs) — à ajuster si le détail des packs
# change, ce n'est pas une nouvelle source de vérité.
_DEFAULT_HIGHLIGHTS: dict[Pack, list[str]] = {
    Pack.starter: [
        "Un site vitrine professionnel, adapté à votre secteur",
        "Surveillance de disponibilité et sauvegardes régulières",
    ],
    Pack.business: [
        "Site complet + révisions illimitées",
        "Présence sur les réseaux sociaux (contenu, calendrier de publication)",
        "Prospection commerciale : fiches qualifiées chaque mois",
    ],
    Pack.premium: [
        "Site complet avec intégrations avancées",
        "Réseaux sociaux multi-plateformes + campagnes",
        "Prospection commerciale à grande échelle",
        "Support prioritaire",
    ],
}


def default_highlights_for_pack(pack: Pack) -> list[str]:
    return _DEFAULT_HIGHLIGHTS[pack]


def generate_offer_pdf(business_name: str, pack: Pack, highlights: list[str]) -> bytes:
    buffer = io.BytesIO()
    doc = canvas.Canvas(buffer, pagesize=A4)
    _width, height = A4

    doc.setFont("Helvetica-Bold", 18)
    doc.drawString(2 * cm, height - 3 * cm, f"Offre de digitalisation — {business_name}")

    doc.setFont("Helvetica", 14)
    doc.drawString(2 * cm, height - 4.5 * cm, f"Pack proposé : {_PACK_LABELS[pack]}")

    doc.setFont("Helvetica", 11)
    y = height - 6 * cm
    for highlight in highlights:
        doc.drawString(2.3 * cm, y, f"• {highlight}")
        y -= 0.8 * cm

    doc.showPage()
    doc.save()
    return buffer.getvalue()


def generate_offer_pptx(business_name: str, pack: Pack, highlights: list[str]) -> bytes:
    """Même contenu que `generate_offer_pdf` (titre, pack, points forts), en une seule
    diapositive — un support PowerPoint sert le même usage commercial (présenter l'offre
    en rendez-vous), pas un deck multi-diapositives à ce stade.

    Layout vide (`slide_layouts[6]`) plutôt qu'un layout à placeholders prédéfinis : les
    index de placeholders varient selon le modèle PowerPoint sous-jacent — des zones de
    texte ajoutées explicitement donnent un contrôle total et reproductible.
    """
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    title_frame = slide.shapes.add_textbox(Cm(1.5), Cm(1), Cm(23), Cm(2)).text_frame
    title_frame.text = f"Offre de digitalisation — {business_name}"
    title_frame.paragraphs[0].font.size = Pt(28)
    title_frame.paragraphs[0].font.bold = True

    pack_frame = slide.shapes.add_textbox(Cm(1.5), Cm(3), Cm(23), Cm(1.5)).text_frame
    pack_frame.text = f"Pack proposé : {_PACK_LABELS[pack]}"
    pack_frame.paragraphs[0].font.size = Pt(20)

    content_frame = slide.shapes.add_textbox(Cm(1.8), Cm(5), Cm(22), Cm(12)).text_frame
    content_frame.word_wrap = True
    for index, highlight in enumerate(highlights):
        paragraph = content_frame.paragraphs[0] if index == 0 else content_frame.add_paragraph()
        paragraph.text = f"• {highlight}"
        paragraph.font.size = Pt(16)

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()
